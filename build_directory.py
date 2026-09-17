"""Rebuild the Bangladesh pet directory (HTML page + Excel) from every source in data/ and apify/results/.

Run after each Apify phase:  python3 build_directory.py
"""
import csv, glob, json, re, collections
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill
from shapely.geometry import shape, Point
from shapely.strtree import STRtree

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "data" / "sources"

D = {r["district"]: r for r in csv.DictReader(open(ROOT / "data" / "districts.csv"))}
NAME_FIX = {"Barisal": "Barishal", "Bogra": "Bogura", "Brahamanbaria": "Brahmanbaria", "Chittagong": "Chattogram",
            "Comilla": "Cumilla", "Jessore": "Jashore", "Maulvibazar": "Moulvibazar", "Nawabganj": "Chapai Nawabganj",
            "Netrakona": "Netrokona", "Dhaka City": "Dhaka", "Savar": "Dhaka"}
geo = json.load(open(ROOT / "data" / "bgd_adm2.geojson"))
POLYS = [shape(f["geometry"]) for f in geo["features"]]
PNAMES = [NAME_FIX.get(f["properties"]["shapeName"], f["properties"]["shapeName"]) for f in geo["features"]]
TREE = STRtree(POLYS)


def district_at(lat, lon):
    p = Point(lon, lat)
    hit = [i for i in TREE.query(p) if POLYS[i].contains(p)]
    return PNAMES[hit[0]] if hit else None


def buckets(text):
    t = text.lower(); b = set()
    if re.search(r"vet|clinic|hospital|doctor|animal care|medicine|ambulance|ভেটেরিনারি|হাসপাতাল|চিকিৎসা", t): b.add("vet")
    if re.search(r"store|shop|mart|food|seller|bazar|market|aquarium|aqua|bird|supplies|শপ", t): b.add("shop")
    if re.search(r"boarding|hotel|hostel|foster|daycare|day care|resort|sitter|sitting|cattery|kennel", t): b.add("hotel")
    if re.search(r"groom|spa\b|salon", t): b.add("groom")
    return sorted(b) or ["other"]


rows = []

# 1. Posha Ghor
for r in json.load(open(SRC / "pg_detail.json")):
    rows.append(dict(name=r["name"].replace("\xa0", " ").replace(" ", " ").strip(), cats=buckets(r["types"]), subtype="",
                     district=r["district"], area=r["area"], address=r["address"], phone=r["phone"], maps=r.get("maps", ""),
                     facebook=r.get("facebook", ""), website=r.get("website", ""), sources=["Posha Ghor"]))
# 2. OpenStreetMap
for r in json.load(open(SRC / "osm_detail.json")):
    rows.append(dict(name=r["name"] if r["name"] != "(unnamed on OSM)" else "", cats=buckets(r["types"]), subtype="",
                     district=r["district"], area=r["area"], address=r["address"], phone=r["phone"], maps=r["maps"],
                     facebook=r["facebook"], website=r["website"], sources=["OpenStreetMap"]))
# 3. Chonk stockists
for r in json.load(open(SRC / "chonk.json")):
    rows.append(dict(name=r["name"].strip(), cats=["shop"], subtype="", district=r["district"], area=r.get("area", ""),
                     address=r.get("address", ""), phone=r.get("phone", ""), maps="", facebook="", website="",
                     sources=["Chonk stockist list"]))
# 4. Facebook pages via Overture
for r in json.load(open(SRC / "overture_pet.json")):
    r = {k: v for k, v in r.items() if k not in ("src_detail", "link")}
    r["sources"] = ["Facebook pages (Overture Maps)"]
    if r.get("subtype", "").startswith("Government"):
        r["cats"] = ["gov"]
    rows.append(r)
# 4b. Government vet hospitals and livestock offices (Department of Livestock Services office lists)
rows.extend(json.load(open(SRC / "dls_offices.json")))

# 5. Google Maps via Apify
PET_RX = re.compile(r"pet|vet|animal|aquarium|groom|kennel|bird shop|cattery|dog|cat food", re.I)
seen_place = set(); g_raw = g_kept = 0
for f in sorted(glob.glob(str(ROOT / "apify" / "results" / "*.json"))):
    for p in json.load(open(f)):
        g_raw += 1
        pid = p.get("placeId") or p.get("url")
        if pid in seen_place:
            continue
        seen_place.add(pid)
        cats_txt = " ".join([p.get("categoryName") or ""] + (p.get("categories") or []))
        if not (PET_RX.search(cats_txt) or PET_RX.search(p.get("title") or "")):
            continue
        if p.get("permanentlyClosed"):
            continue
        loc = p.get("location") or {}
        dist = district_at(loc["lat"], loc["lng"]) if loc.get("lat") else None
        if not dist:
            continue
        g_kept += 1
        hours = "; ".join(f"{h.get('day')}: {h.get('hours')}" for h in (p.get("openingHours") or [])).replace(" ", " ")
        site = p.get("website") or ""
        fb = site if "facebook.com" in site else ""
        site = "" if fb else site
        rows.append(dict(name=(p.get("title") or "").strip(), cats=buckets(cats_txt + " " + (p.get("title") or "")),
                         subtype=p.get("categoryName") or "", district=dist, area=p.get("neighborhood") or "",
                         address=p.get("address") or "", phone=p.get("phoneUnformatted") or p.get("phone") or "",
                         maps=p.get("url") or "", facebook=fb, website=site,
                         rating=p.get("totalScore"), reviews=p.get("reviewsCount") or 0, hours=hours,
                         temp_closed=bool(p.get("temporarilyClosed")), sources=["Google Maps"]))

# Merge duplicates: same name + district, or same phone number
out, by_key, by_phone = [], {}, {}
for r in rows:
    r["district"] = NAME_FIX.get(r["district"], r["district"])
    if r["district"] not in D:
        r["district"] = ""
    r["division"] = D[r["district"]]["division"] if r["district"] else ""
    for k in ("subtype", "rating", "hours"):
        r.setdefault(k, None if k == "rating" else "")
    r.setdefault("reviews", 0); r.setdefault("conf", None); r.setdefault("temp_closed", False)
    key = re.sub(r"[^a-z0-9ঀ-৿]", "", r["name"].lower())
    if not key:
        r["name"] = "Unnamed place"; out.append(r); continue
    ph = re.sub(r"\D", "", r.get("phone") or "")[-10:]
    hit = by_key.get((key, r["district"])) or (by_phone.get(ph) if len(ph) == 10 else None)
    if hit:
        for f in ("address", "phone", "maps", "facebook", "website", "area", "subtype", "hours", "district", "division"):
            if not hit[f] and r[f]:
                hit[f] = r[f]
        if r["rating"] is not None:
            hit["rating"], hit["reviews"] = r["rating"], r["reviews"]
            if r["maps"]: hit["maps"] = r["maps"]
        hit["cats"] = sorted(set(hit["cats"]) | set(r["cats"]))
        hit["sources"] = sorted(set(hit["sources"]) | set(r["sources"]))
        continue
    by_key[(key, r["district"])] = r
    out.append(r)
    if len(ph) == 10:
        by_phone.setdefault(ph, r)

CATS = ["shop", "vet", "hotel", "groom", "other", "gov"]
dists = []
for d, m in D.items():
    allrs = [r for r in out if r["district"] == d]
    rs = [r for r in allrs if r["cats"] != ["gov"]]   # private businesses only; govt offices counted in "gov"
    dists.append(dict(district=d, division=m["division"], pop=int(m["population_2022"]), area=int(m["area_km2"]),
                      total=len(rs), reviews=sum(r["reviews"] for r in rs),
                      **{c: sum(1 for r in allrs if c in r["cats"]) for c in CATS}))

meta = dict(counts={s: sum(1 for r in out if s in r["sources"]) for s in sorted({s for r in out for s in r["sources"]})},
            google_raw=g_raw, google_kept=g_kept)
data = dict(rows=out, districts=dists, meta=meta)

# HTML
tpl = (ROOT / "data" / "page_template.html").read_text()
page = tpl.replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))
(ROOT / "bangladesh-pet-directory.html").write_text(page)
# Standalone copy for static hosting (Vercel serves index.html at /)
(ROOT / "index.html").write_text('<!doctype html>\n<html lang="en">\n<head>\n<meta charset="utf-8">\n'
    '<meta name="viewport" content="width=device-width, initial-scale=1, viewport-fit=cover">\n'
    '<meta name="description" content="Directory of pet shops, vets, pet hotels, groomers and government vet hospitals across all 64 districts of Bangladesh.">\n'
    '<link rel="icon" href="data:image/svg+xml,<svg xmlns=%22http://www.w3.org/2000/svg%22 viewBox=%220 0 100 100%22><text y=%22.9em%22 font-size=%2290%22>🐾</text></svg>">\n'
    '</head>\n<body>\n' + page + '\n</body>\n</html>\n')

# Excel
L = {"shop": "Pet shop / food", "vet": "Pet hospital / vet (private)", "hotel": "Pet hotel / boarding", "groom": "Grooming", "other": "Other", "gov": "Govt vet hospital / livestock office"}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Directory"
ws.append(["Name", "Category", "Type", "Division", "District", "Area", "Address", "Phone", "Google rating", "Google reviews",
           "Opening hours", "Map", "Facebook", "Website", "Sources", "Confidence (Facebook data)"])
for r in sorted(out, key=lambda r: (r["division"] or "zz", r["district"] or "zz", -r["reviews"], r["name"])):
    ws.append([r["name"], ", ".join(L[c] for c in r["cats"]), r["subtype"], r["division"], r["district"], r["area"], r["address"],
               r["phone"], r["rating"], r["reviews"] or None, r["hours"], r["maps"], r["facebook"], r["website"],
               " + ".join(r["sources"]), r["conf"]])
ds = wb.create_sheet("64 districts")
ds.append(["District", "Division", "Population 2022", "Known businesses", "Pet shops / food", "Vets", "Hotels / boarding",
           "Grooming", "Other", "Govt vet offices", "People per known business", "Google reviews (all businesses)"])
for x in sorted(dists, key=lambda x: -x["total"]):
    ds.append([x["district"], x["division"], x["pop"], x["total"], x["shop"], x["vet"], x["hotel"], x["groom"], x["other"], x["gov"],
               round(x["pop"] / x["total"]) if x["total"] else None, x["reviews"]])
for sh, widths in ((ws, [36, 26, 22, 12, 14, 16, 48, 16, 10, 10, 40, 30, 34, 30, 34, 12]), (ds, [18, 12, 16, 16, 16, 8, 16, 10, 8, 16, 22, 22])):
    for c in sh[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F6A64")
    for i, w in enumerate(widths):
        sh.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    sh.freeze_panes = "A2"; sh.auto_filter.ref = sh.dimensions
wb.save(ROOT / "bangladesh_pet_business_directory.xlsx")

print(f"{len(out)} unique businesses | sources: {meta['counts']} | Google places read {g_raw}, pet-related kept {g_kept}")
print("Districts with none:", [x["district"] for x in dists if x["total"] == 0])
